"""
CareerCast Resume Parser

Provides functionality for:
- Resume text extraction
- Skill extraction
- Education extraction
- TXT parsing
- PDF parsing
- DOCX parsing
- Complete resume parsing
"""

import os
import re


# ============================================================
# KNOWN SKILLS
# ============================================================

KNOWN_SKILLS = [
    # Programming Languages
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "go",
    "rust",
    "php",
    "ruby",
    "kotlin",
    "swift",

    # Web Development
    "html",
    "css",
    "react",
    "angular",
    "vue",
    "node.js",
    "node",
    "express",
    "django",
    "flask",
    "fastapi",

    # Data Science / Machine Learning
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "ai",
    "data science",
    "data analysis",
    "data analytics",
    "statistics",
    "nlp",
    "natural language processing",
    "computer vision",
    "tensorflow",
    "pytorch",
    "scikit-learn",
    "sklearn",
    "keras",
    "xgboost",
    "random forest",
    "pandas",
    "numpy",
    "matplotlib",
    "seaborn",

    # Databases
    "sql",
    "mysql",
    "postgresql",
    "postgres",
    "mongodb",
    "oracle",
    "sqlite",
    "redis",

    # Cloud
    "aws",
    "amazon web services",
    "azure",
    "google cloud",
    "gcp",
    "docker",
    "kubernetes",

    # DevOps
    "git",
    "github",
    "gitlab",
    "jenkins",
    "ci/cd",
    "linux",
    "unix",

    # Big Data
    "hadoop",
    "spark",
    "apache spark",
    "kafka",

    # Tools / Technologies
    "tableau",
    "power bi",
    "excel",
    "jira",
    "mlflow",
    "streamlit",
    "selenium",
    "rest api",
    "api",

    # Software Engineering
    "object oriented programming",
    "oop",
    "data structures",
    "algorithms",
    "software development",
    "software engineering",

    # Business / Management
    "project management",
    "product management",
    "business analysis",
    "communication",
    "leadership",
    "problem solving",
]


# ============================================================
# EDUCATION KEYWORDS
# ============================================================

EDUCATION_KEYWORDS = [
    "bachelor",
    "bachelor's",
    "bachelors",
    "master",
    "master's",
    "masters",
    "phd",
    "doctorate",
    "diploma",
    "degree",
    "b.tech",
    "btech",
    "m.tech",
    "mtech",
    "b.e",
    "be",
    "m.e",
    "me",
    "b.sc",
    "bsc",
    "m.sc",
    "msc",
    "bca",
    "mca",
    "mba",
    "computer science",
    "information technology",
    "engineering",
]


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(text):
    """
    Extract known skills from resume text.

    Parameters
    ----------
    text : str
        Resume text.

    Returns
    -------
    list
        List of detected skills.
    """

    if not text:
        return []

    text_lower = text.lower()

    found_skills = []

    for skill in KNOWN_SKILLS:
        # Escape skill names so characters such as + and . are
        # treated literally.
        escaped_skill = re.escape(skill)

        # Word-boundary matching for reliable extraction.
        pattern = rf"(?<!\w){escaped_skill}(?!\w)"

        if re.search(pattern, text_lower):
            found_skills.append(skill)

    # Remove duplicates while preserving order.
    return list(dict.fromkeys(found_skills))


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

def extract_education(text):
    """
    Extract education information from resume text.

    Parameters
    ----------
    text : str
        Resume text.

    Returns
    -------
    list
        Detected education qualifications.
    """

    if not text:
        return []

    text_lower = text.lower()

    education = []

    education_patterns = {
        "Bachelor's Degree": [
            "bachelor",
            "bachelor's",
            "bachelors",
            "b.tech",
            "btech",
            "b.e",
            "be",
            "b.sc",
            "bsc",
            "bca",
        ],
        "Master's Degree": [
            "master",
            "master's",
            "masters",
            "m.tech",
            "mtech",
            "m.e",
            "me",
            "m.sc",
            "msc",
            "mca",
            "mba",
        ],
        "PhD": [
            "phd",
            "ph.d",
            "doctorate",
        ],
        "Diploma": [
            "diploma",
        ],
    }

    for education_name, patterns in education_patterns.items():
        for pattern in patterns:
            if pattern.lower() in text_lower:
                education.append(education_name)
                break

    # If no specific degree was found, check generic education
    # keywords.
    if not education:
        for keyword in EDUCATION_KEYWORDS:
            if keyword.lower() in text_lower:
                education.append(keyword.title())
                break

    return list(dict.fromkeys(education))


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt(file):
    """
    Extract text from a TXT file.

    Supports:
    - Normal Python file objects
    - Flask UploadedFile objects
    """

    if hasattr(file, "read"):
        content = file.read()

        # Bytes → string
        if isinstance(content, bytes):
            try:
                return content.decode("utf-8")
            except UnicodeDecodeError:
                return content.decode("latin-1")

        return str(content)

    # Fallback when a file path is provided.
    if isinstance(file, (str, os.PathLike)):
        with open(file, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    raise ValueError("Invalid TXT file object.")


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf(file):
    """
    Extract text from a PDF file.

    Requires PyPDF2.
    """

    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise ImportError(
            "PyPDF2 is required to read PDF resumes. "
            "Install it with: pip install PyPDF2"
        ) from exc

    # File path
    if isinstance(file, (str, os.PathLike)):
        reader = PdfReader(str(file))

    # File object
    else:
        # Reset stream when possible.
        try:
            file.seek(0)
        except Exception:
            pass

        reader = PdfReader(file)

    text_parts = []

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text_parts.append(page_text)

    return "\n".join(text_parts)


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx(file):
    """
    Extract text from a DOCX file.

    Requires python-docx.
    """

    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(
            "python-docx is required to read DOCX resumes. "
            "Install it with: pip install python-docx"
        ) from exc

    # File path
    if isinstance(file, (str, os.PathLike)):
        document = Document(str(file))

    # File object
    else:
        try:
            file.seek(0)
        except Exception:
            pass

        document = Document(file)

    text_parts = []

    # Paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_resume_text(file):
    """
    Extract raw text from a PDF, DOCX, or TXT resume.

    Supports both Flask-style uploaded files and normal
    Python file objects.

    Flask example:
        uploaded_file.filename

    Normal Python file example:
        open("resume.txt", "rb")
        file.name
    """

    # --------------------------------------------------------
    # Determine filename
    # --------------------------------------------------------

    filename = getattr(file, "filename", "")

    # Normal Python file objects generally use .name.
    if not filename:
        filename = getattr(file, "name", "")

    # Path-like input
    if not filename and isinstance(file, (str, os.PathLike)):
        filename = os.fspath(file)

    filename = str(filename).strip().lower()

    # --------------------------------------------------------
    # TXT FILE
    # --------------------------------------------------------

    if filename.endswith(".txt"):
        return extract_txt(file)

    # --------------------------------------------------------
    # PDF FILE
    # --------------------------------------------------------

    elif filename.endswith(".pdf"):
        return extract_pdf(file)

    # --------------------------------------------------------
    # DOCX FILE
    # --------------------------------------------------------

    elif filename.endswith(".docx"):
        return extract_docx(file)

    # --------------------------------------------------------
    # UNSUPPORTED FILE
    # --------------------------------------------------------

    else:
        raise Exception(
            "Unsupported file format. "
            "Please upload PDF, DOCX or TXT."
        )


# ============================================================
# COMPLETE RESUME PARSING
# ============================================================

def parse_resume_text(text):
    """
    Perform complete parsing of resume text.

    Parameters
    ----------
    text : str
        Raw resume text.

    Returns
    -------
    dict
        Parsed resume information.
    """

    if not text:
        text = ""

    text = str(text)

    skills = extract_skills(text)
    education = extract_education(text)

    return {
        "skills": skills,
        "education": education,
        "text_length": len(text),
    }


# ============================================================
# COMPLETE FILE PARSING
# ============================================================

# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_txt(file):
    """
    Extract text from a TXT file.

    Supports:
    - UTF-8
    - UTF-8 with BOM
    - UTF-16
    - UTF-16 with BOM
    - Latin-1
    - Normal Python file objects
    - Flask UploadedFile objects
    - File paths
    """

    # --------------------------------------------------------
    # Read raw bytes
    # --------------------------------------------------------

    if isinstance(file, (str, os.PathLike)):
        with open(file, "rb") as f:
            content = f.read()

    elif hasattr(file, "read"):
        try:
            file.seek(0)
        except Exception:
            pass

        content = file.read()

    else:
        raise ValueError("Invalid TXT file object.")

    # --------------------------------------------------------
    # Already decoded string
    # --------------------------------------------------------

    if isinstance(content, str):
        return content

    if not content:
        return ""

    # --------------------------------------------------------
    # Detect Byte Order Marks
    # --------------------------------------------------------

    # UTF-16 Little Endian
    if content.startswith(b"\xff\xfe"):
        return content.decode("utf-16")

    # UTF-16 Big Endian
    if content.startswith(b"\xfe\xff"):
        return content.decode("utf-16")

    # UTF-8 with BOM
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")

    # --------------------------------------------------------
    # Try UTF-8
    # --------------------------------------------------------

    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        pass

    # --------------------------------------------------------
    # Try UTF-16
    # --------------------------------------------------------

    try:
        return content.decode("utf-16")
    except UnicodeDecodeError:
        pass

    # --------------------------------------------------------
    # Final fallback
    # --------------------------------------------------------

    return content.decode("latin-1")